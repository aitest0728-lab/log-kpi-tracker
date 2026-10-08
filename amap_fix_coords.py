#!/usr/bin/env python3
"""Correct the suspicious map coordinates with AMap (高德) and report every change in an Excel file.

  python amap_fix_coords.py --test            # 1) try 3 estates, print what AMap returns, write NOTHING
  python amap_fix_coords.py                   # 2) geocode every row of the list, update index.html, write the summary Excel
  python amap_fix_coords.py --dry-run         # same as 2) but index.html is left untouched (Excel is still written)
  python amap_fix_coords.py --from-json       # re-apply a previous run to index.html without calling AMap again

Defaults (override with options):
  --env   C:\\Users\\chipanl\\Downloads\\Amap.env
  --xlsx  Adjusted_Coordinates_Map_Check.xlsx   --html index.html   --out Coordinates_Adjustment_Summary.xlsx

Engines (--engine auto|rest|browser, default auto):
  rest     AMap Web-Service REST API (restapi.amap.com). Needs a key created as type "Web服务".
  browser  AMap JS API. Works with a "Web端(JS API)" key + its securityJsCode. The script starts a tiny local web page
           (http://127.0.0.1:8765), opens it in your normal browser, and the page does the geocoding and hands the results back.
           Keep that browser tab open and visible until the console says "done".
  auto     try REST first; if AMap answers INVALID_USER_KEY (a JS key used on REST), switch to browser.

Search rule: the CORRECT lat/long is found from the estate NAME only (Chinese, then English). The old / wrong lat/long is never
used to search, reverse-geocode or filter AMap - it is only reported in the Excel (Old Latitude / Longitude, Moved km).

A result is accepted only if it is inside Hong Kong AND within the row's 'Zone limit (km)' of the delivery-zone reference
centre, and is not just a district/city centroid - so same-name places elsewhere (彩虹, 九龍, 順利 ...) are rejected.
Accepted rows go into the window.__MAP_FIX__ block of index.html (method 'amap'); a backup of index.html is made first.

Needs: pip install pandas openpyxl requests
"""
import argparse, datetime as dt, json, math, os, re, shutil, sys, threading, time, webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

HK = (22.13, 22.57, 113.82, 114.45)
COARSE = {"国家", "省", "市", "区县"}            # AMap 'level' values that are only a centroid, never a building
DEFAULT_ENV = r"C:\Users\chipanl\Downloads\Amap.env"


def in_hk(la, ln): return HK[0] <= la <= HK[1] and HK[2] <= ln <= HK[3]
def km(a, b, c, d): return math.hypot((a - c) * 111, (b - d) * 111 * math.cos(math.radians(a)))


# ---------- GCJ-02 -> WGS-84 (only used with --gcj2wgs) ----------
def gcj_to_wgs(lat, lng):
    a, ee = 6378245.0, 0.00669342162296594323
    def tl(x, y):
        r = -100 + 2*x + 3*y + .2*y*y + .1*x*y + .2*math.sqrt(abs(x))
        r += (20*math.sin(6*x*math.pi) + 20*math.sin(2*x*math.pi)) * 2/3
        r += (20*math.sin(y*math.pi) + 40*math.sin(y/3*math.pi)) * 2/3
        return r + (160*math.sin(y/12*math.pi) + 320*math.sin(y*math.pi/30)) * 2/3
    def tn(x, y):
        r = 300 + x + 2*y + .1*x*x + .1*x*y + .1*math.sqrt(abs(x))
        r += (20*math.sin(6*x*math.pi) + 20*math.sin(2*x*math.pi)) * 2/3
        r += (20*math.sin(x*math.pi) + 40*math.sin(x/3*math.pi)) * 2/3
        return r + (150*math.sin(x/12*math.pi) + 300*math.sin(x/30*math.pi)) * 2/3
    dl, dn = tl(lng - 105, lat - 35), tn(lng - 105, lat - 35)
    rl = lat / 180 * math.pi; mg = 1 - ee*math.sin(rl)**2; sm = math.sqrt(mg)
    dl = dl * 180 / ((a*(1-ee)) / (mg*sm) * math.pi); dn = dn * 180 / (a / sm * math.cos(rl) * math.pi)
    return lat - dl, lng - dn


# ---------- Amap.env ----------
def load_env(path):
    txt = open(path, encoding="utf-8-sig").read()
    d = {}
    try:
        j = json.loads(txt)
        if isinstance(j, dict): d = {str(k): str(v) for k, v in j.items()}
    except Exception: pass
    if not d:
        lines = [l.strip() for l in txt.splitlines()]
        for i, line in enumerate(lines):
            line = re.sub(r"^\s*export\s+", "", line)
            if not line or line.startswith(("#", "//")): continue
            m = re.match(r"^([A-Za-z_][\w.\- ]*?)\s*[=:]\s*(.*)$", line)
            if m:
                v = m.group(2).strip().strip("\"' ,;.")
                if not v and i + 1 < len(lines) and lines[i + 1] and not re.match(r"^[A-Za-z_][\w.\- ]*?\s*[=:]", lines[i + 1]):
                    v = lines[i + 1].rstrip(",;").strip("\"'")      # value written on the next line
                d[m.group(1).strip().replace(" ", "_")] = v
    d["_RAW"] = txt
    return d


def pick_keys(env):
    raw = env.get("_RAW", "")
    it = {k.upper(): v for k, v in env.items() if v and k != "_RAW"}
    sec = next((v for k, v in it.items() if re.search(r"SECURITY|JSCODE|SECRET", k)), None)
    keys = {k: v for k, v in it.items() if "KEY" in k and not re.search(r"SECURITY|SECRET", k)}
    rest = next((v for k, v in keys.items() if re.search(r"REST|WEB_?SERVICE|SERVER", k)), None)
    js = next((v for k, v in keys.items() if re.search(r"JS|WEB_?API|BROWSER", k)), None)
    gen = next(iter(keys.values()), None)
    if not gen:      # key not labelled: look for key=xxxx in a URL, else any 32-character hex token that is not the security code
        m = re.search(r"[?&]key=([0-9a-fA-F]{32})", raw)
        hexes = [h for h in dict.fromkeys(re.findall(r"\b[0-9a-fA-F]{32}\b", raw)) if h != sec]
        gen = m.group(1) if m else (hexes[0] if len(hexes) == 1 else None)
    return {"rest": rest or gen, "js": js or gen, "sec": sec}


def mask(s): return f"{s[:4]}…{s[-3:]} (len {len(s)})" if s and len(s) > 8 else "(none)"


# ---------- reading the list ----------
def find(df, *pats, ban=()):
    for c in df.columns:
        n = str(c).lower()
        if all(re.search(p, n) for p in pats) and not any(re.search(b, n) for b in ban): return c


def load_index_data(html_text):
    m = re.search(r"window\.__MAP_DATA__=(\{.*?\});/\*MAP_DATA_END\*/", html_text, re.S)
    out = {}
    if not m: return out
    D = json.loads(m.group(1))
    for ar in D["areas"]:
        for r in ar["e"]:
            out[r[0]] = {"en": r[1], "zh": r[2], "dd": D["dd"][r[3]] if r[3] < len(D["dd"]) else "", "zone": D["zones"][r[5]] if r[5] < len(D["zones"]) else "",
                         "area": ar["k"], "lat": r[6] if len(r) > 7 else None, "lng": r[7] if len(r) > 7 else None}
    return out


def read_rows(xlsx, idx, limit=None):
    import pandas as pd
    try: df = pd.read_excel(xlsx, sheet_name="Adjusted List")
    except Exception: df = pd.read_excel(xlsx, sheet_name=0)
    cc = dict(code=find(df, "estate", "code"), en=find(df, "name", r"\ben\b|\(en"), zh=find(df, "name", r"\bzh\b|\(zh"),
              zlat=find(df, "zone", "lat"), zlng=find(df, "zone", r"lon|lng"), lim=find(df, "limit"),
              olat=find(df, "lat", ban=("verified", "zone", "ref")), olng=find(df, r"lon|lng", ban=("verified", "zone", "ref")))
    if not cc["code"]: sys.exit(f"No 'Estate Code' column in {xlsx}. Columns found: {list(df.columns)}")
    num = lambda v: float(v) if v is not None and v == v and str(v).strip() != "" else None
    rows = []
    for _, r in df.iterrows():
        code = str(r[cc["code"]]).strip()
        if not code or code.lower() == "nan": continue
        g = lambda k: (r[cc[k]] if cc[k] else None)
        info = idx.get(code, {})
        zc = (num(g("zlat")), num(g("zlng")))
        rows.append(dict(code=code, en=str(g("en") if g("en") == g("en") and g("en") is not None else info.get("en", "")),
                         zh=str(g("zh") if g("zh") == g("zh") and g("zh") is not None else info.get("zh", "")),
                         zc=zc if None not in zc else None, limit=num(g("lim")) or 2.0,
                         olat=num(g("olat")) if num(g("olat")) is not None else info.get("lat"),
                         olng=num(g("olng")) if num(g("olng")) is not None else info.get("lng"),
                         dd=info.get("dd", ""), area=info.get("area", ""), zone=info.get("zone", "")))
    if limit: rows = rows[:limit]
    if not cc["zlat"]: print("WARNING: no 'Zone ref. Latitude' column - only the inside-Hong-Kong check can be applied.")
    return rows


# ---------- query variants ----------
NOTE = re.compile(r"只限|交收|內部使用|deliver", re.I)
def strip_notes(s):
    s = re.sub(r"[\u200b-\u200f\ufeff]", "", s or "")
    return re.sub(r"\s+", " ", re.sub(r"[(（]([^)）]*)[)）]", lambda m: " " if NOTE.search(m.group(1)) else m.group(0), s)).strip()


def variants(zh, en, n):
    zh, en = strip_notes(zh), strip_notes(en)
    no_par = re.sub(r"\s+", " ", re.sub(r"[(（][^)）]*[)）]", " ", zh)).strip()
    addr = " ".join(re.findall(r"[(（]([^)）]*(?:[號街路道])[^)）]*)[)）]", zh))
    base_zh = re.split(r"[\s\-–—/／]|(?<=[\u4e00-\u9fff])\d", no_par)[0]
    en_clean = re.sub(r"\s+", " ", re.sub(r"[(（][^)）]*[)）]", " ", en)).strip()
    base_en = re.split(r"\s+(?:BLOCK|TOWER|PHASE|PHRASE|HOUSE)\b|\s+-\s*", en_clean, flags=re.I)[0].strip(" -")
    out = []
    for v in (no_par, (no_par.split(" ")[0] + " " + addr).strip() if addr else "", base_zh, en_clean, base_en):
        if v and len(v) >= 2 and v not in out: out.append(v)
    return out[:n]


# ---------- candidates / selection ----------
def choose(cands, zc, limit, wgs):
    for c in cands:
        if "err" in c: continue
        la, ln = c["lat"], c["lng"]
        if wgs: la, ln = gcj_to_wgs(la, ln)
        if not in_hk(la, ln) or c.get("level") in COARSE: continue
        d = km(la, ln, *zc) if zc else None
        if zc and d > limit: continue
        return {**c, "lat": la, "lng": ln, "dist": d}


def nearest(cands, zc):
    best = None
    for c in cands:
        if "err" in c or not in_hk(c["lat"], c["lng"]) or c.get("level") in COARSE: continue
        d = km(c["lat"], c["lng"], *zc) if zc else 0
        if best is None or d < best[0]: best = (d, c)
    return best


class AmapError(Exception):
    def __init__(self, info, code=""): super().__init__(f"{info} ({code})"); self.info, self.code = info, code
FATAL = ("INVALID_USER_KEY", "USERKEY_PLAT_NOMATCH", "INVALID_USER_SCODE", "INVALID_USER_DOMAIN", "INSUFFICIENT_PRIVILEGES", "DAILY_QUERY_OVER_LIMIT")


# ---------- engine 1: REST ----------
class Rest:
    name = "rest"
    def __init__(self, key, base="https://restapi.amap.com"):
        import requests
        self.key, self.base, self.s, self.calls = key, base.rstrip("/"), requests.Session(), 0
    def _get(self, path, **p):
        for a in range(4):
            self.calls += 1
            j = self.s.get(self.base + path, params={**p, "key": self.key}, timeout=30).json()
            if str(j.get("status")) == "1": return j
            info = str(j.get("info", ""))
            if "QPS" in info or j.get("infocode") in ("10014", "10019", "10020"): time.sleep(1 + a); continue
            raise AmapError(info, j.get("infocode", ""))
        raise AmapError("rate limited", "10014")
    @staticmethod
    def _pt(loc):
        ln, la = loc.split(","); return float(la), float(ln)
    def run(self, row, queries, zc, limit, wgs):
        cands = []
        for q in queries:
            for kind in ("poi", "geo"):
                if kind == "poi":
                    j = self._get("/v3/place/text", keywords=q, city="香港", citylimit="true", offset=10, page=1)
                    new = [dict(zip(("lat", "lng"), self._pt(p["location"])), name=p.get("name", ""), addr=p.get("address") if isinstance(p.get("address"), str) else "",
                                level="", src="poi", q=q) for p in j.get("pois", []) if p.get("location")]
                else:
                    j = self._get("/v3/geocode/geo", address=q, city="香港")
                    new = [dict(zip(("lat", "lng"), self._pt(g["location"])), name=g.get("formatted_address", ""), addr="", level=g.get("level", ""),
                                src="geocode", q=q) for g in j.get("geocodes", []) if g.get("location")]
                cands += new
                time.sleep(0.05)
                if choose(new, zc, limit, wgs): return cands
        return cands


# ---------- engine 2: JS API through a local page ----------
PAGE = r"""<!doctype html><meta charset="utf-8"><title>AMap geocoding bridge</title>
<body style="font:14px system-ui;padding:20px;max-width:900px"><h3>AMap geocoding &mdash; keep this tab open and visible</h3><div id="st">starting…</div><pre id="log" style="color:#555"></pre>
__SEC__
<script src="https://webapi.amap.com/maps?v=2.0&key=__KEY__" onerror="window.__amapErr=1"></script>
<script>
const say=m=>{post('/log',{m}).catch(()=>{});$('log').textContent+=m+'\n'};
const $=id=>document.getElementById(id), log=m=>{$('log').textContent+=m+'\n'}, sleep=ms=>new Promise(r=>setTimeout(r,ms));
const post=(p,o)=>fetch(p,{method:'POST',body:JSON.stringify(o)});
const HK=[22.13,22.57,113.82,114.45], COARSE=['国家','省','市','区县'], DELAY=__DELAY__;
const inHK=(a,b)=>a>=HK[0]&&a<=HK[1]&&b>=HK[2]&&b<=HK[3];
const km=(a,b,c,d)=>Math.hypot((a-c)*111,(b-d)*111*Math.cos(a*Math.PI/180));
const ll=p=>[p.getLat?p.getLat():p.lat, p.getLng?p.getLng():p.lng];
let geocoder, ps, calls=0, to=0;
function call(fn){return new Promise(res=>{let n=0,q=0,id=0;const go=()=>{const my=++id;const t=setTimeout(()=>{if(my!==id)return;if(q<1){q++;log('no answer from AMap after 15s - retrying');go()}else{id++;res({s:'error',r:'timeout (no answer from AMap in 15s)'})}},15000);
  fn((s,r)=>{if(my!==id)return;clearTimeout(t);calls++;if(s==='error'&&n<3&&/QPS/i.test(JSON.stringify(r))){n++;id++;setTimeout(go,900*n)}else{id++;res({s,r})}})};go()})}
async function poi(q){const {s,r}=await call(cb=>ps.search(q,cb));
  if(s==='error')return{c:[],err:JSON.stringify(r)};
  return{c:((r&&r.poiList&&r.poiList.pois)||[]).filter(p=>p.location).map(p=>{const [a,b]=ll(p.location);return{lat:a,lng:b,name:p.name||'',addr:typeof p.address==='string'?p.address:'',level:'',src:'poi'}})}}
async function geo(q){const {s,r}=await call(cb=>geocoder.getLocation(q,cb));
  if(s==='error')return{c:[],err:JSON.stringify(r)};
  return{c:((r&&r.geocodes)||[]).filter(g=>g.location).map(g=>{const [a,b]=ll(g.location);return{lat:a,lng:b,name:g.formattedAddress||'',addr:'',level:g.level||'',src:'geocode'}})}}
const ok=(x,t)=>inHK(x.lat,x.lng)&&!COARSE.includes(x.level)&&(!t.zc||km(x.lat,x.lng,t.zc[0],t.zc[1])<=t.limit);
async function main(){
  for(let n=0;n<40&&typeof AMap==='undefined'&&!window.__amapErr;n++)await sleep(250);
  if(typeof AMap==='undefined'){const m='AMap JS file did not load in this browser ('+(window.__amapErr?'the request to webapi.amap.com failed - blocked by network/firewall/ad-blocker?':'timeout')+'). Try opening https://webapi.amap.com/maps?v=2.0 in this browser.';await post('/fatal',{msg:m});$('st').textContent=m;return}
  say('AMap script loaded, loading Geocoder/PlaceSearch plugins...');
  const pl=await Promise.race([new Promise(r=>{try{AMap.plugin(['AMap.Geocoder','AMap.PlaceSearch'],()=>r('ok'))}catch(e){r('error: '+e)}}),sleep(20000).then(()=>'timeout')]);
  if(pl!=='ok'){const m='AMap plugins did not load ('+pl+'). The AMap map script loaded but its plugin files could not be fetched - usually a key / securityJsCode problem (check the key is type Web端(JS API) and the security code belongs to this key) or a blocker.';await post('/fatal',{msg:m});$('st').textContent=m;return}
  say('plugins ready - starting estates');
  geocoder=new AMap.Geocoder({city:'香港'}); ps=new AMap.PlaceSearch({city:'香港',citylimit:true,pageSize:10,pageIndex:1});
  const tasks=await (await fetch('/tasks')).json(); let i=0;
  for(const t of tasks){
    const cands=[]; let done=false; $('st').textContent=(i+1)+' / '+tasks.length+' estates ('+calls+' AMap requests so far) - now: '+t.code+(document.hidden?'  (TAB IS HIDDEN - the browser may pause it, bring it to the front)':'');
    for(const q of t.queries){ for(const kind of ['poi','geo']){
      const r=await (kind==='poi'?poi(q):geo(q)); await sleep(DELAY);
      if(i===0&&!window.__f1){window.__f1=1;say('first reply ('+kind+', "'+q+'"): '+(r.err?'ERROR '+r.err.slice(0,250):r.c.length+' hits'))}
      if(r.err&&/^"?timeout/.test(r.err)){ if(++to>=3){const m='AMap stopped answering ('+to+' requests in a row timed out after 15s). Likely causes: the key hit its daily quota / is being throttled, or the browser/network is blocking webapi.amap.com or restapi.amap.com. Progress so far is saved - wait a while (or check the quota in the AMap console) and rerun with --resume.';await post('/fatal',{msg:m});$('st').textContent=m;return} } else to=0;
      if(r.err){ if(/INVALID_USER|USERKEY|SCODE|DOMAIN|PRIVILEGES|OVER_LIMIT/.test(r.err)){await post('/fatal',{msg:r.err});$('st').textContent='AMap rejected the request: '+r.err;return}
                 cands.push({err:r.err,q,src:kind}) }
      r.c.forEach(x=>cands.push({...x,q}));
      if(r.c.some(x=>ok(x,t))){done=true;break}
    } if(done)break }
    await post('/result',{code:t.code,cands,calls}); i++; $('st').textContent=i+' / '+tasks.length+' estates'; if(i%20==0)log(i+' done')
  }
  await post('/done',{calls}); $('st').textContent='done - you can close this tab';
}
main().catch(async e=>{await post('/fatal',{msg:String(e)});$('st').textContent='error: '+e});
</script>"""


class Browser:
    name = "browser"
    def __init__(self, key, sec, port, timeout, delay=0.35):
        self.key, self.sec, self.port, self.timeout, self.delay, self.calls = key, sec, port, timeout, delay, 0
        self.preload, self.partial = {}, None
    def run_all(self, tasks):
        results, state = dict(self.preload), {"done": False, "fatal": None}
        total = len(tasks); tasks = [t for t in tasks if t["code"] not in results]
        if self.preload: print(f"Resuming: {len(results)} estates already done, {len(tasks)} left.")
        def save():
            if self.partial: json.dump(results, open(self.partial, "w", encoding="utf-8"), ensure_ascii=False)
        sec = f'<script>window._AMapSecurityConfig={{securityJsCode:{json.dumps(self.sec)}}};</script>' if self.sec else ""   # must be its own tag, BEFORE the AMap script
        page = PAGE.replace("__SEC__", sec).replace("__KEY__", self.key).replace("__DELAY__", str(int(self.delay * 1000)))
        class H(BaseHTTPRequestHandler):
            def log_message(self, *a): pass
            def _send(self, body, ctype):
                b = body.encode("utf-8"); self.send_response(200); self.send_header("Content-Type", ctype); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)
            def do_GET(self):
                if self.path.startswith("/tasks"): self._send(json.dumps(tasks, ensure_ascii=False), "application/json; charset=utf-8")
                else: self._send(page, "text/html; charset=utf-8")
            def do_POST(self):
                o = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or b"{}")
                if self.path == "/result":
                    results[o["code"]] = o["cands"]; state["calls"] = o.get("calls", 0); state["t"] = time.time()
                    if len(results) % 5 == 0: save()
                elif self.path == "/log": print("  [page] " + o.get("m", ""))
                elif self.path == "/done": state["done"] = True; state["calls"] = o.get("calls", 0)
                elif self.path == "/fatal": state["fatal"] = o.get("msg", "unknown")
                self._send("{}", "application/json")
        srv = ThreadingHTTPServer(("127.0.0.1", self.port), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        url = f"http://127.0.0.1:{self.port}/"
        print(f"Opening {url} in your browser - keep that tab open and visible until this window says done.")
        webbrowser.open(url)
        t0, last, warned = time.time(), -1, 0
        state["t"] = t0
        try:
            while not state["done"] and not state["fatal"] and time.time() - t0 < self.timeout:
                time.sleep(1)
                if len(results) != last and len(results) % 10 == 0: last = len(results); print(f"  {len(results)}/{total} estates  ({state.get('calls', 0)} AMap requests so far)")
                if time.time() - state["t"] > 90 and time.time() - warned > 90:
                    warned = time.time(); print(f"  ...no new estate for {int(time.time() - state['t'])}s ({len(results)}/{total}). Look at the browser tab: is it visible, and does it show an error? "
                                                f"Ctrl+C is safe - progress is saved; rerun with --resume.")
        except KeyboardInterrupt:
            save(); srv.shutdown(); sys.exit(f"Stopped by you. {len(results)}/{total} estates saved to {self.partial} - continue with: python amap_fix_coords.py --resume")
        save(); srv.shutdown(); self.calls = state.get("calls", 0)
        if state["fatal"]: raise AmapError(state["fatal"])
        if not state["done"]: print(f"WARNING: timed out after {self.timeout}s with {len(results)}/{len(tasks)} estates answered.")
        return results


# ---------- output ----------
def write_excel(path, rows, res, engine, calls, html_note, wgs):
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    adj, na = [], []
    for r in rows:
        x = res[r["code"]]
        base = {"Estate Code": r["code"], "Estate Name (EN)": r["en"], "Estate Name (ZH)": r["zh"], "District": r["dd"], "Area": r["area"], "Zone": r["zone"]}
        if x["pick"]:
            p = x["pick"]; shift = km(r["olat"], r["olng"], p["lat"], p["lng"]) if r["olat"] is not None else None
            adj.append({**base, "Old Latitude": r["olat"], "Old Longitude": r["olng"], "New Latitude": round(p["lat"], 6), "New Longitude": round(p["lng"], 6),
                        "Moved (km)": round(shift, 2) if shift is not None else None, "New point to zone centre (km)": round(p["dist"], 2) if p["dist"] is not None else None,
                        "Zone limit (km)": r["limit"], "AMap source": p["src"], "AMap matched name": p["name"], "AMap address": p["addr"], "AMap level": p.get("level", ""), "Query used": p["q"]})
        else:
            na.append({**base, "Old Latitude": r["olat"], "Old Longitude": r["olng"], "Zone limit (km)": r["limit"], "Reason": x["reason"],
                       "Nearest AMap hit": x["near"], "Queries tried": " | ".join(x["queries"])})
    A, N = pd.DataFrame(adj), pd.DataFrame(na)
    full = []
    for r in rows:
        x = res[r["code"]]; p = x["pick"]
        full.append({"Estate Code": r["code"], "Estate Name (EN)": r["en"], "Estate Name (ZH)": r["zh"], "District": r["dd"], "Area": r["area"], "Zone": r["zone"],
                     "Status": "Adjusted" if p else "Not adjusted", "Old Latitude": r["olat"], "Old Longitude": r["olng"],
                     "New Latitude": round(p["lat"], 6) if p else None, "New Longitude": round(p["lng"], 6) if p else None,
                     "Moved (km)": round(km(r["olat"], r["olng"], p["lat"], p["lng"]), 2) if p and r["olat"] is not None else None,
                     "AMap matched name": p["name"] if p else "", "AMap address": p["addr"] if p else "", "Note": "" if p else x["reason"]})
    F = pd.DataFrame(full)
    moved = A["Moved (km)"].dropna() if len(A) else []
    summ = pd.DataFrame([
        ("Run time", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")), ("AMap engine", engine), ("Rows checked", len(rows)),
        ("Adjusted (AMap, inside Hong Kong and zone limit)", len(adj)), ("Not adjusted (kept as before)", len(na)),
        ("Average move (km)", round(float(moved.mean()), 2) if len(moved) else ""), ("Largest move (km)", round(float(moved.max()), 2) if len(moved) else ""),
        ("AMap API requests", calls), ("GCJ-02 -> WGS-84 conversion applied", "yes" if wgs else "no"), ("index.html", html_note)], columns=["Item", "Value"])
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in (("Summary", summ), ("Full List", F), ("Adjusted", A), ("Not Adjusted", N)):
            df.to_excel(w, sheet_name=name, index=False); ws = w.sheets[name]
            for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F172A"); c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if name != "Summary" and ws.max_row > 1: ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col[:200]) + 2))


def apply_html(html_path, fixes):
    s = open(html_path, encoding="utf-8").read()
    m = re.search(r"/\*MAP_FIX_START\*/window\.__MAP_FIX__=(\{.*?\});/\*MAP_FIX_END\*/", s, re.S)
    if not m: sys.exit(f"{html_path} has no MAP_FIX block - it is not the updated index.html (v37.0). Nothing changed.")
    bak = f"{html_path}.bak-{dt.datetime.now():%Y%m%d-%H%M%S}"; shutil.copy2(html_path, bak)
    fix = json.loads(m.group(1)); fix.update(fixes)
    block = "/*MAP_FIX_START*/window.__MAP_FIX__=" + json.dumps(fix, ensure_ascii=False) + ";/*MAP_FIX_END*/"
    open(html_path, "w", encoding="utf-8").write(s[:m.start()] + block + s[m.end():])
    return bak, len(fix)


def diagnose(a):
    import requests
    env = load_env(a.env) if os.path.exists(a.env) else {}; k = pick_keys(env)
    key, sec = a.key or k["js"], a.sec or k["sec"]
    print(f"key {mask(key)}  (hex chars only: {bool(key and re.fullmatch(r'[0-9a-fA-F]+', key))})   security code {mask(sec)}")
    for label, url, par in (("REST geocode", "https://restapi.amap.com/v3/geocode/geo", {"address": "珍寶廣場", "city": "香港", "key": key}),
                            ("JS loader", "https://webapi.amap.com/maps", {"v": "2.0", "key": key})):
        try:
            r = requests.get(url, params=par, timeout=20); print(f"{label}: HTTP {r.status_code}, {len(r.content)} bytes\n   {r.text[:300]!r}")
        except Exception as e: print(f"{label}: FAILED from Python -> {e}")
    print("\nIf 'JS loader' works here but the browser tab fails, the browser (ad-blocker / proxy) is blocking webapi.amap.com.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--env", default=DEFAULT_ENV); ap.add_argument("--xlsx", default="Adjusted_Coordinates_Map_Check.xlsx"); ap.add_argument("--html", default="index.html")
    ap.add_argument("--out", default="Coordinates_Adjustment_Summary.xlsx"); ap.add_argument("--json", default="amap_geocode_results.json")
    ap.add_argument("--engine", choices=("auto", "rest", "browser"), default="auto"); ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--timeout", type=int, default=3600, help="browser engine: seconds to wait"); ap.add_argument("--max-queries", type=int, default=4)
    ap.add_argument("--limit", type=int, help="only the first N rows (trial run)"); ap.add_argument("--test", action="store_true")
    ap.add_argument("--resume", action="store_true", help="browser engine: continue from amap_partial.json instead of starting again"); ap.add_argument("--partial", default="amap_partial.json")
    ap.add_argument("--dry-run", action="store_true"); ap.add_argument("--diagnose", action="store_true", help="check the key / network and print raw AMap answers"); ap.add_argument("--from-json", action="store_true")
    ap.add_argument("--gcj2wgs", action="store_true", help="convert AMap's GCJ-02 answers to WGS-84 (normally NOT needed in Hong Kong)")
    ap.add_argument("--key", help="AMap key (overrides the env file)"); ap.add_argument("--sec", help="securityJsCode (overrides the env file)")
    ap.add_argument("--rest-base", default="https://restapi.amap.com", help=argparse.SUPPRESS)
    a = ap.parse_args()

    if a.from_json:
        fixes = json.load(open(a.json, encoding="utf-8"))["fixes"]
        bak, n = apply_html(a.html, fixes); print(f"Re-applied {len(fixes)} fixes -> {a.html} (total overrides {n}); backup {bak}"); return

    if a.diagnose: return diagnose(a); 
    html_text = open(a.html, encoding="utf-8").read()
    rows = read_rows(a.xlsx, load_index_data(html_text), 3 if a.test else a.limit)
    if not os.path.exists(a.env): sys.exit(f"Cannot find {a.env}. Use --env to point at your Amap.env file.")
    env = load_env(a.env); keys = pick_keys(env)
    if a.key: keys["rest"] = keys["js"] = a.key
    if a.sec: keys["sec"] = a.sec
    print(f"Amap.env: variables {sorted(k for k in env if k != '_RAW')}  | REST key {mask(keys['rest'])}  JS key {mask(keys['js'])}  security code {mask(keys['sec'])}")
    if not keys["rest"] and not keys["js"]: sys.exit("No key found in the env file (expected a variable with KEY in its name).")
    tasks = [{"code": r["code"], "queries": variants(r["zh"], r["en"], a.max_queries), "zc": r["zc"], "limit": r["limit"]} for r in rows]

    cand, engine, calls = {}, None, 0
    if a.engine in ("auto", "rest") and keys["rest"]:
        try:
            eng = Rest(keys["rest"], a.rest_base); engine = "rest"
            for i, (r, t) in enumerate(zip(rows, tasks), 1):
                cand[r["code"]] = eng.run(r, t["queries"], r["zc"], r["limit"], a.gcj2wgs)
                if i % 20 == 0: print(f"  {i}/{len(rows)} rows")
            calls = eng.calls
        except AmapError as e:
            if a.engine == "rest" or not (keys["js"] and e.info in ("INVALID_USER_KEY", "USERKEY_PLAT_NOMATCH", "INVALID_USER_DOMAIN", "SERVICE_NOT_AVAILABLE")):
                sys.exit(f"AMap REST refused the request: {e}\n  -> this key is probably a JS-API key. Run with --engine browser, or create a 'Web服务' key in the AMap console.")
            print(f"AMap REST said {e} - this is a JS-API key, switching to the browser engine."); cand, engine = {}, None
    if engine is None:
        if not keys["js"]: sys.exit("No usable key for the browser engine.")
        try:
            br = Browser(keys["js"], keys["sec"], a.port, a.timeout); br.partial = a.partial
            if a.resume and os.path.exists(a.partial):
                br.preload = json.load(open(a.partial, encoding="utf-8"))
                bad = [k for k, v in br.preload.items() if any("err" in c for c in v)]
                for k in bad: del br.preload[k]          # estates that hit an AMap error / timeout are searched again
                if bad: print(f"  {len(bad)} saved estates had AMap errors and will be searched again.")
            cand = br.run_all(tasks); engine = "browser"; calls = br.calls
        except AmapError as e: sys.exit(f"AMap JS API refused the request: {e}")

    res, fixes = {}, {}
    for r, t in zip(rows, tasks):
        cs = cand.get(r["code"], [])
        pick = choose(cs, r["zc"], r["limit"], a.gcj2wgs)
        errs = [c["err"] for c in cs if "err" in c]
        nb = nearest(cs, r["zc"]) if cs else None
        reason = ("AMap returned no usable result" if not [c for c in cs if "err" not in c] else
                  f"all {len([c for c in cs if 'err' not in c])} AMap hits were outside Hong Kong, outside the {r['limit']:g} km zone limit, or only a district centroid")
        if errs and not pick: reason += f" | AMap errors: {errs[0][:120]}"
        res[r["code"]] = {"pick": pick, "reason": reason, "queries": t["queries"],
                          "near": (f"{nb[1]['name']} ({nb[0]:.1f} km from zone centre)" if nb and r["zc"] else (nb[1]["name"] if nb else ""))}
        if pick: fixes[r["code"]] = [round(pick["lat"], 6), round(pick["lng"], 6), "amap", f"AMap {pick['src']}: {pick['name']}"[:90]]
    ok = len(fixes)
    print(f"\nAMap ({engine}): {ok}/{len(rows)} rows have a verified coordinate; {len(rows) - ok} stay unchanged.")
    for r in rows[:3 if a.test else 0]:
        p = res[r["code"]]["pick"]; print(f"  {r['code']} {r['en']}: " + (f"{p['lat']:.6f}, {p['lng']:.6f}  <- {p['src']} '{p['name']}' ({p['dist']:.2f} km from zone centre)" if p else res[r["code"]]["reason"]))
    if a.test: print("\n--test: nothing written. If the lines above look right, run again without --test."); return

    html_note = "not modified (--dry-run)"
    if not a.dry_run and fixes:
        bak, n = apply_html(a.html, fixes); html_note = f"updated with {ok} fixes (total overrides {n}); backup {os.path.basename(bak)}"
    json.dump({"engine": engine, "fixes": fixes}, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    write_excel(a.out, rows, res, engine, calls, html_note, a.gcj2wgs)
    print(f"{html_note}\nSummary: {a.out}   Fixes JSON: {a.json}  (re-apply later with --from-json)\nReload the dashboard; Data Check should shrink by about {ok}.")


if __name__ == "__main__": main()
